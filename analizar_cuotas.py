import openpyxl
import pandas as pd

# Leer archivo Excel
excel_file = "CUOTAS 2025 INSM vigente para cristian (2).xlsx"

# Obtener nombres de hojas
wb = openpyxl.load_workbook(excel_file, data_only=True)
print(f"Hojas disponibles: {wb.sheetnames}\n")

# Leer cada hoja
for sheet_name in wb.sheetnames:
    print(f"\n{'='*80}")
    print(f"HOJA: {sheet_name}")
    print(f"{'='*80}")
    
    # Leer con pandas
    df = pd.read_excel(excel_file, sheet_name=sheet_name)
    
    print(f"Filas: {len(df)}, Columnas: {df.shape[1]}")
    print(f"Columnas: {df.columns.tolist()}\n")
    
    # Mostrar primeras filas
    print("Primeras 3 filas:")
    print(df.head(3).to_string())
    print("\n")
