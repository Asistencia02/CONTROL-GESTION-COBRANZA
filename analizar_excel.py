#!/usr/bin/env python3
# -*- coding: utf-8 -*-
import sys
import io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

import openpyxl
from collections import defaultdict

xlsx_path = r"C:\Users\isipp1206\Downloads\programa contable\gestion-cobranzas\CUOTAS 2025 INSM vigente para cristian.xlsx"

wb = openpyxl.load_workbook(xlsx_path)

mapeo = {
    "INICIAL2026": (2, 4),
    "PRIMARIA2026": (2, 5),
    "SECUNDARIA2026": (2, 6),
    "HIGIENE2026": (1, 3),
    "ANALISTA2026": (1, 1)
}

dnis_por_inst = defaultdict(list)
duplicados_multi_hoja = {}

print("="*80)
print("ANALISIS DE EXCEL - BUSQUEDA DE DUPLICADOS Y ERRORES")
print("="*80 + "\n")

for sheet_name in wb.sheetnames:
    if sheet_name not in mapeo:
        print(f"Hoja ignorada: {sheet_name}\n")
        continue
        
    inst_id, carrera_id = mapeo[sheet_name]
    ws = wb[sheet_name]
    
    print(f"Hoja: {sheet_name} (Institucion {inst_id}, Carrera {carrera_id})")
    print(f"   Filas: {ws.max_row}, Columnas: {ws.max_column}")
    
    # Encontrar índices
    headers = {}
    for col in range(1, ws.max_column + 1):
        h = str(ws.cell(1, col).value or "").upper().strip()
        if h in ["DNI", "APELLIDO", "NOMBRES", "INSCRIPCION"]:
            headers[h] = col - 1
    
    if "DNI" not in headers:
        print(f"   ERROR: COLUMNA DNI NO ENCONTRADA\n")
        continue
    
    dnis_hoja = {}
    duplicados_internos = []
    
    for row in range(2, ws.max_row + 1):
        dni_raw = str(ws.cell(row, headers["DNI"] + 1).value or "").replace(".", "").strip()
        apellido = str(ws.cell(row, headers.get("APELLIDO", 1) + 1).value or "").strip() if "APELLIDO" in headers else ""
        nombre = str(ws.cell(row, headers.get("NOMBRES", 1) + 1).value or "").strip() if "NOMBRES" in headers else ""
        
        if dni_raw and dni_raw != "0":
            if dni_raw in dnis_hoja:
                duplicados_internos.append({"dni": dni_raw, "fila_1": dnis_hoja[dni_raw]["fila"], "fila_2": row})
            else:
                dnis_hoja[dni_raw] = {"fila": row, "apellido": apellido, "nombre": nombre}
    
    # Check cross-sheet
    for dni, data in dnis_hoja.items():
        clave = f"INST{inst_id}_{dni}"
        if clave in dnis_por_inst:
            duplicados_multi_hoja[clave] = duplicados_multi_hoja.get(clave, []) + [{
                "hoja": sheet_name,
                "fila": data["fila"],
                "apellido": data["apellido"],
                "nombre": data["nombre"]
            }]
        else:
            dnis_por_inst[clave] = [{
                "hoja": sheet_name,
                "fila": data["fila"],
                "apellido": data["apellido"],
                "nombre": data["nombre"]
            }]
    
    print(f"   OK Estudiantes unicos: {len(dnis_hoja)}")
    if duplicados_internos:
        print(f"   ATENCION! DUPLICADOS EN MISMA HOJA: {len(duplicados_internos)}")
        for dup in duplicados_internos[:3]:
            print(f"      - DNI {dup['dni']}: filas {dup['fila_1']} y {dup['fila_2']}")
    print()

print("\n" + "="*80)
print("DUPLICADOS ENTRE HOJAS (MISMO ESTUDIANTE EN MULTIPLES CARRERAS)")
print("="*80 + "\n")

if duplicados_multi_hoja:
    for clave, registros in sorted(duplicados_multi_hoja.items()):
        print(f"ERROR! {clave}:")
        for reg in registros:
            print(f"   - Hoja: {reg['hoja']}, Fila: {reg['fila']}, {reg['apellido']}, {reg['nombre']}")
    print()
else:
    print("OK NO HAY DUPLICADOS ENTRE HOJAS\n")

print(f"\nRESUMEN:")
print(f"   Total DNIs unicos: {len(dnis_por_inst)}")
print(f"   DNIs duplicados (multiples carreras): {len(duplicados_multi_hoja)}")
